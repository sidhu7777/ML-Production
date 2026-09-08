"""
PPT Automation Module
Provides PPTAutomator class to manipulate PowerPoint presentations.
Preserves shape dimensions, font styles, and layout structure.
"""

import os
import io
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# ─────────────────────────────────────────────────────────────────
# LEGEND PNG GENERATOR
# ─────────────────────────────────────────────────────────────────

def generate_legend_png(items, title="", output_path=None):
    """
    Generate a crisp legend PNG with colored circle bullet markers
    matching the reference PPT template style.

    Parameters
    ----------
    items       : list of dicts with:
                  - 'text': pre-formatted label string, OR
                  - 'min', 'max', 'count', 'pct', 'label', 'color'
                  - 'color': hex color string
    title       : str - optional legend title
    output_path : str - path to save the PNG
    """
    from PIL import Image, ImageDraw, ImageFont

    if not items:
        return None

    scale = 3
    pt_to_px = (96.0 / 72.0) * scale

    font_size_pt = 7.5
    title_size_pt = 8.0
    row_h_pt = 13.5
    pad_left_pt = 5.0
    pad_right_pt = 8.0
    pad_y_pt = 5.0
    dot_r_pt = 2.8
    dot_pad_pt = 4.0

    f_size = int(font_size_pt * pt_to_px)
    t_size = int(title_size_pt * pt_to_px)
    r_h = int(row_h_pt * pt_to_px)
    p_left = int(pad_left_pt * pt_to_px)
    p_right = int(pad_right_pt * pt_to_px)
    p_y = int(pad_y_pt * pt_to_px)
    dot_r = int(dot_r_pt * pt_to_px)
    dot_pad = int(dot_pad_pt * pt_to_px)
    dot_col_w = dot_r * 2

    font_path = r"C:\Windows\Fonts\segoeui.ttf"
    font = None
    title_font = None
    for fp in [font_path, r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\calibri.ttf"]:
        if os.path.exists(fp):
            try:
                font = ImageFont.truetype(fp, f_size)
                title_font = ImageFont.truetype(fp, t_size)
                break
            except Exception:
                pass
    if font is None:
        font = ImageFont.load_default()
        title_font = ImageFont.load_default()

    rows = []
    for idx, it in enumerate(items):
        color = it.get("color", "#999999")
        if "text" in it and it["text"]:
            text = str(it["text"])
        else:
            min_v = it.get("min")
            max_v = it.get("max")
            count = it.get("count", 0)
            pct = it.get("pct", 0.0)
            label = it.get("label", "")
            is_first = (idx == 0)
            is_last = (idx == len(items) - 1)
            has_numeric = (
                min_v is not None and max_v is not None
                and str(min_v).strip() != "" and str(max_v).strip() != ""
            )
            if has_numeric:
                try:
                    f_min = float(min_v)
                    f_max = float(max_v)
                    if is_first and not is_last:
                        text = "Below {:.2f} ({}) {:.1f}%".format(f_max, count, pct)
                    elif is_last and not is_first:
                        text = "Above {:.2f} ({}) {:.1f}%".format(f_min, count, pct)
                    else:
                        text = ">= {:.2f} to < {:.2f} ({}) {:.1f}%".format(f_min, f_max, count, pct)
                except (ValueError, TypeError):
                    text = "{} ({}) {:.1f}%".format(label or min_v, count, pct)
            else:
                lbl = label or it.get("name") or str(min_v or max_v or "---")
                text = "{} ({}) {:.1f}%".format(lbl, count, pct)
        rows.append({"text": text, "color": color})

    dummy = Image.new("RGBA", (10, 10))
    draw_dummy = ImageDraw.Draw(dummy)
    max_text_w = 0
    for r in rows:
        bbox = draw_dummy.textbbox((0, 0), r["text"], font=font)
        w = bbox[2] - bbox[0]
        if w > max_text_w:
            max_text_w = w

    title_str = str(title).strip() if title else ""
    title_h = 0
    if title_str:
        t_bbox = draw_dummy.textbbox((0, 0), title_str, font=title_font)
        t_w = t_bbox[2] - t_bbox[0]
        if t_w > max_text_w:
            max_text_w = t_w
        title_h = int(14.0 * pt_to_px)

    total_w = p_left + dot_col_w + dot_pad + max_text_w + p_right
    total_h = p_y * 2 + title_h + len(rows) * r_h

    img = Image.new("RGBA", (total_w, total_h), (255, 255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.rectangle([0, 0, total_w - 1, total_h - 1], outline=(0, 0, 0), width=int(1 * scale))

    curr_y = p_y
    if title_str:
        draw.text((p_left, curr_y), title_str, fill=(0, 0, 0), font=title_font)
        curr_y += title_h

    for r in rows:
        color_hex = str(r["color"]).lstrip("#")
        try:
            rgb = tuple(int(color_hex[i:i+2], 16) for i in (0, 2, 4))
        except Exception:
            rgb = (128, 128, 128)

        mid_y = curr_y + r_h // 2
        cx = p_left + dot_r
        cy = mid_y
        draw.ellipse(
            [cx - dot_r, cy - dot_r, cx + dot_r, cy + dot_r],
            fill=rgb + (255,),
            outline=(50, 50, 50),
        )

        ty = curr_y + (r_h - int(f_size * 1.15)) // 2
        draw.text((p_left + dot_col_w + dot_pad, ty), r["text"], fill=(0, 0, 0), font=font)
        curr_y += r_h

    if output_path:
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        img.save(output_path, "PNG")

    w_pts = total_w / pt_to_px
    h_pts = total_h / pt_to_px
    w_emu = int(w_pts * 12700)
    h_emu = int(h_pts * 12700)
    return w_emu, h_emu


# ─────────────────────────────────────────────────────────────────
# PPT AUTOMATOR
# ─────────────────────────────────────────────────────────────────

class PPTAutomator:
    def __init__(self, template_path):
        if not os.path.exists(template_path):
            raise FileNotFoundError(f"Template not found: {template_path}")
        self.template_path = template_path
        self.prs = Presentation(template_path)

    def find_slide_by_title_or_index(self, target):
        if isinstance(target, int):
            if 1 <= target <= len(self.prs.slides):
                return self.prs.slides[target - 1]
            raise IndexError(f"Slide index {target} out of range")
        target_str = str(target).strip().lower()
        for slide in self.prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame and target_str in shape.text_frame.text.strip().lower():
                    return slide
        raise ValueError(f"No slide found matching '{target}'")

    def update_cover_metadata(self, test_type=None, region=None, area=None, completion_date=None, total_km=None):
        slide1 = self.prs.slides[0]
        target_shape = None
        for shape in slide1.shapes:
            if shape.has_text_frame and "Test Type" in shape.text_frame.text:
                target_shape = shape
                break
        if not target_shape:
            return

        if region and str(region).strip().upper().startswith(("POLYGON", "MULTIPOLYGON", "GEOMETRY")):
            region = "SEO"
        if area and str(area).strip().upper().startswith(("POLYGON", "MULTIPOLYGON", "GEOMETRY")):
            area = "Test Site"

        field_map = {}
        if test_type:        field_map["Test Type"] = f"Test Type : {test_type}"
        if region:           field_map["Region"]    = f"Region : {region}"
        if area:             field_map["Area"]      = f"Area : {area}"
        if completion_date:  field_map["DT Completion Date"] = f"DT Completion Date (YYYY/MM/DD) : {completion_date}"
        if total_km is not None:
            if isinstance(total_km, (int, float)):
                field_map["Total kilometers"] = f"Total kilometers : {total_km:.2f} km"
            else:
                s = str(total_km).strip()
                field_map["Total kilometers"] = s if s.startswith("Total kilometers") else f"Total kilometers : {s}"

        tf = target_shape.text_frame
        for p in tf.paragraphs:
            p_text = p.text.strip()
            for prefix, new_val in field_map.items():
                if p_text.startswith(prefix) or prefix in p_text:
                    if p.runs:
                        p.runs[0].text = new_val
                        for r in p.runs[1:]:
                            r.text = ""
                    else:
                        p.text = new_val
                    break

    def update_slide_text(self, slide_target, replacements: dict):
        """Replace text across all shapes on a slide."""
        try:
            slide = self.find_slide_by_title_or_index(slide_target)
        except (IndexError, ValueError):
            return
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for r in para.runs:
                    for old, new_val in replacements.items():
                        if old in r.text:
                            if new_val is None:
                                r.text = ""
                            else:
                                r.text = str(new_val)

    def replace_slide_image(self, slide_target, new_image_path, target_type="main"):
        """
        Replace a picture shape on a slide and remove any srcRect cropping
        that was baked into the template.

        target_type='main'   -> largest picture (the map)
        target_type='legend' -> second-largest picture (the legend)
        """
        if not os.path.exists(new_image_path):
            print(f"  [PPT | Slide {slide_target} | {target_type}] SKIP — image file not found: {new_image_path}")
            return False
        try:
            slide = self.find_slide_by_title_or_index(slide_target)
        except (IndexError, ValueError) as e:
            print(f"  [PPT | Slide {slide_target} | {target_type}] SKIP — slide not found: {e}")
            return False

        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        pic_count = len(pictures)

        if not pictures:
            print(f"  [PPT | Slide {slide_target} | {target_type}] SKIP — slide has 0 picture shapes (cannot inject)")
            return False

        pictures.sort(key=lambda s: s.width * s.height, reverse=True)

        if target_type.lower() == "main":
            target_pic = pictures[0]
        elif target_type.lower() == "legend":
            if pic_count < 2:
                try:
                    slide.shapes.add_picture(
                        new_image_path,
                        left=6057900,
                        top=3508502,
                        width=1270000,
                    )
                    print(
                        f"  [PPT | Slide {slide_target} | legend] OK — added new legend picture shape "
                        f"'{os.path.basename(new_image_path)}'"
                    )
                    return True
                except Exception as e:
                    print(
                        f"  [PPT | Slide {slide_target} | legend] FAIL — could not add legend shape: {e}"
                    )
                    return False
            target_pic = pictures[1]
        else:
            try:
                target_pic = pictures[int(target_type)]
            except (ValueError, IndexError):
                target_pic = pictures[0]

        blip_nodes = target_pic._element.xpath(".//a:blip")
        if not blip_nodes:
            print(f"  [PPT | Slide {slide_target} | {target_type}] FAIL — no blip node found in picture shape")
            return False

        # Add new image part to the slide part to ensure this slide has an isolated relationship.
        # This prevents mutating shared template media (such as image19.png shared by slides 11, 13, 14, 16).
        image_part, new_rId = slide.part.get_or_add_image_part(new_image_path)
        blip_nodes[0].attrib[
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
        ] = new_rId

        # Remove any hardcoded srcRect crop tags left over from template screenshots
        for src_rect in target_pic._element.xpath(".//a:srcRect"):
            src_rect.getparent().remove(src_rect)

        print(
            f"  [PPT | Slide {slide_target} | {target_type}] OK — "
            f"injected '{os.path.basename(new_image_path)}' "
            f"(slide has {pic_count} picture shape(s))"
        )
        return True

    def remove_legend_image(self, slide_target):
        """Remove secondary picture shapes (e.g. template legends) from a slide."""
        try:
            slide = self.find_slide_by_title_or_index(slide_target)
        except (IndexError, ValueError):
            return False
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if len(pictures) >= 2:
            pictures.sort(key=lambda s: s.width * s.height, reverse=True)
            for pic in pictures[1:]:
                try:
                    pic._element.getparent().remove(pic._element)
                except Exception:
                    pass
            return True
        return False

    def remove_text_shapes_by_content(self, slide_target, search_text):
        """Remove text-box shapes whose text contains search_text (case-insensitive)."""
        try:
            slide = self.find_slide_by_title_or_index(slide_target)
        except (IndexError, ValueError):
            return False
        removed = 0
        for shape in list(slide.shapes):
            if shape.has_text_frame and search_text.lower() in shape.text_frame.text.strip().lower():
                shape._element.getparent().remove(shape._element)
                removed += 1
        return removed > 0

    def update_table_cell(self, slide_target, table_index, row_idx, col_idx, new_value):
        try:
            slide = self.find_slide_by_title_or_index(slide_target)
        except (IndexError, ValueError):
            return False
        tables = [s for s in slide.shapes if s.has_table]
        if not tables or table_index >= len(tables):
            return False
        tbl = tables[table_index].table
        if 0 <= row_idx < len(tbl.rows) and 0 <= col_idx < len(tbl.columns):
            cell = tbl.cell(row_idx, col_idx)
            p = cell.text_frame.paragraphs[0] if cell.text_frame.paragraphs else None
            if p:
                if p.runs:
                    p.runs[0].text = str(new_value)
                    for r in p.runs[1:]:
                        r.text = ""
                else:
                    p.text = str(new_value)
            else:
                cell.text = str(new_value)
            return True
        return False

    def update_table_row_by_key(self, slide_target, table_index, key_string, new_value, value_col_idx=2):
        try:
            slide = self.find_slide_by_title_or_index(slide_target)
        except (IndexError, ValueError):
            return False
        tables = [s for s in slide.shapes if s.has_table]
        if not tables or table_index >= len(tables):
            return False
        tbl = tables[table_index].table
        for row_idx, row in enumerate(tbl.rows):
            cell_texts = [row.cells[i].text.strip() for i in range(min(value_col_idx, len(row.cells)))]
            if key_string.lower() in " ".join(cell_texts).lower():
                self.update_table_cell(slide_target, table_index, row_idx, value_col_idx, new_value)
                return True
        return False

    def save(self, output_path):
        """Save the presentation, freeing PowerPoint lock if open so output_path is reliably overwritten."""
        import time
        import subprocess
        out_dir = os.path.dirname(os.path.abspath(output_path))
        if out_dir and not os.path.exists(out_dir):
            os.makedirs(out_dir, exist_ok=True)

        try:
            self.prs.save(output_path)
            print(f"[PPT Pipeline] SUCCESS! Saved -> {output_path}")
            return output_path
        except PermissionError:
            print(f"[Notice] '{os.path.basename(output_path)}' is open in PowerPoint. Terminating PowerPoint to overwrite target...")
            try:
                subprocess.run(
                    ["powershell", "-Command", "Stop-Process -Name POWERPNT -Force -ErrorAction SilentlyContinue"],
                    capture_output=True, timeout=5
                )
                time.sleep(1.0)
                self.prs.save(output_path)
                print(f"[PPT Pipeline] SUCCESS! Overwrote -> {output_path}")
                return output_path
            except Exception as e:
                print(f"[Notice] Could not release PowerPoint lock: {e}. Trying fallback filenames...")

        base, ext = os.path.splitext(output_path)
        candidates = [f"{base}_new{ext}", f"{base}_{int(time.time())}{ext}"]
        for cand in candidates:
            try:
                self.prs.save(cand)
                print(f"[PPT Pipeline] SUCCESS! Saved -> {cand}")
                return cand
            except PermissionError:
                print(f"[Notice] '{os.path.basename(cand)}' is also locked. Trying next fallback...")
        raise PermissionError(f"Could not save presentation to any candidate path in {out_dir}")
